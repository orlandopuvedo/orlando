"""Build a tailored resume (structured + text) from the Master Profile — reorder/select, never invent."""
from __future__ import annotations

import io

from app.ai.base import AIProvider
from app.services.generation import tailored_summary
from app.services.profile_service import profile_skill_terms
from app.services.text_utils import contains_term, normalize_text


def _relevance(text: str, job_norm: str) -> int:
    words = {w for w in normalize_text(text).split() if len(w) > 3}
    return sum(1 for w in words if contains_term(job_norm, w))


def build_tailored_resume(profile, job=None, focus: str | None = None, provider: AIProvider | None = None,
                          max_highlights: int = 5) -> tuple[dict, list[str], str]:
    """Returns (content, warnings, generated_by)."""
    job_text = "\n".join(x for x in [job.title, job.description, job.requirements] if x) if job else (focus or "")
    job_norm = normalize_text(job_text)
    warnings: list[str] = []
    generated_by = "template"

    skills = profile_skill_terms(profile)
    matched = [s for s in skills if contains_term(job_norm, s)]
    ordered_skills = matched + [s for s in skills if s not in matched]

    experience = []
    for exp in profile.work_experience or []:
        hl = list(exp.get("highlights") or [])
        hl.sort(key=lambda h: _relevance(h, job_norm), reverse=True)
        experience.append({**exp, "highlights": hl[:max_highlights]})

    projects = sorted(profile.projects or [], key=lambda p: _relevance(f"{p.get('name')} {p.get('description')}", job_norm), reverse=True)[:4]
    achievements = sorted(profile.achievements or [], key=lambda a: _relevance(a if isinstance(a, str) else str(a), job_norm), reverse=True)[:5]

    summary = profile.professional_summary or ""
    if job is not None:
        gen = tailored_summary(profile, job, provider)
        if gen.text:
            summary = gen.text
            warnings.extend(gen.warnings)
            generated_by = gen.generated_by

    content = {
        "header": {
            "full_name": profile.full_name, "headline": profile.headline,
            "email": profile.email, "phone": profile.phone, "location": profile.location,
            "linkedin_url": profile.linkedin_url, "portfolio_url": profile.portfolio_url,
        },
        "summary": summary,
        "skills": ordered_skills[:30],
        "highlighted_skills": matched[:12],
        "experience": experience,
        "education": profile.education or [],
        "certifications": profile.certifications or [],
        "languages": profile.languages or [],
        "projects": projects,
        "achievements": achievements,
    }
    if not profile.work_experience:
        warnings.append("Your profile has no work experience yet — add it in Profile before applying.")
    return content, warnings, generated_by


def _date_range(e: dict) -> str:
    start = e.get("start_date") or ""
    end = "Present" if e.get("current") else (e.get("end_date") or "")
    return f"{start} – {end}".strip(" –")


def render_text(content: dict, markdown: bool = False) -> str:
    h = content.get("header", {})
    out: list[str] = []
    H = (lambda s: f"## {s}") if markdown else (lambda s: s.upper())
    out.append((f"# {h.get('full_name') or ''}") if markdown else (h.get("full_name") or ""))
    if h.get("headline"):
        out.append(h["headline"])
    contact = " | ".join(x for x in [h.get("email"), h.get("phone"), h.get("location"), h.get("linkedin_url"), h.get("portfolio_url")] if x)
    if contact:
        out.append(contact)
    if content.get("summary"):
        out += ["", H("Summary"), content["summary"]]
    if content.get("skills"):
        out += ["", H("Skills"), ", ".join(content["skills"])]
    if content.get("experience"):
        out += ["", H("Experience")]
        for e in content["experience"]:
            line = " — ".join(x for x in [e.get("title"), e.get("company"), e.get("location")] if x)
            out.append((f"### {line}" if markdown else line) + (f"  ({_date_range(e)})" if _date_range(e) else ""))
            if e.get("description"):
                out.append(e["description"])
            for hl in e.get("highlights") or []:
                out.append(f"- {hl}")
            out.append("")
    if content.get("projects"):
        out += [H("Projects")]
        for p in content["projects"]:
            out.append(f"- {p.get('name')}: {p.get('description') or ''}".rstrip(": "))
        out.append("")
    if content.get("education"):
        out += [H("Education")]
        for ed in content["education"]:
            out.append(" — ".join(x for x in [ed.get("degree"), ed.get("field"), ed.get("institution"), ed.get("end_date")] if x))
        out.append("")
    if content.get("certifications"):
        out += [H("Certifications")]
        for c in content["certifications"]:
            out.append("- " + " — ".join(x for x in [c.get("name"), c.get("issuer"), c.get("date")] if x) if isinstance(c, dict) else f"- {c}")
        out.append("")
    if content.get("languages"):
        out += [H("Languages")]
        out.append(", ".join(
            (f"{lang.get('language')} ({lang.get('proficiency')})" if lang.get("proficiency") else lang.get("language"))
            if isinstance(lang, dict) else str(lang) for lang in content["languages"]))
    return "\n".join(out).strip() + "\n"


def render_docx(content: dict) -> bytes:
    import docx
    from docx.shared import Pt

    d = docx.Document()
    style = d.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)
    h = content.get("header", {})
    d.add_heading(h.get("full_name") or "Resume", level=0)
    if h.get("headline"):
        d.add_paragraph(h["headline"])
    contact = " | ".join(x for x in [h.get("email"), h.get("phone"), h.get("location"), h.get("linkedin_url"), h.get("portfolio_url")] if x)
    if contact:
        d.add_paragraph(contact)
    if content.get("summary"):
        d.add_heading("Summary", level=1)
        d.add_paragraph(content["summary"])
    if content.get("skills"):
        d.add_heading("Skills", level=1)
        d.add_paragraph(", ".join(content["skills"]))
    if content.get("experience"):
        d.add_heading("Experience", level=1)
        for e in content["experience"]:
            p = d.add_paragraph()
            p.add_run(" — ".join(x for x in [e.get("title"), e.get("company")] if x)).bold = True
            if _date_range(e):
                p.add_run(f"   {_date_range(e)}")
            if e.get("description"):
                d.add_paragraph(e["description"])
            for hl in e.get("highlights") or []:
                d.add_paragraph(hl, style="List Bullet")
    if content.get("projects"):
        d.add_heading("Projects", level=1)
        for pr in content["projects"]:
            d.add_paragraph(f"{pr.get('name')}: {pr.get('description') or ''}".rstrip(": "), style="List Bullet")
    if content.get("education"):
        d.add_heading("Education", level=1)
        for ed in content["education"]:
            d.add_paragraph(" — ".join(x for x in [ed.get("degree"), ed.get("field"), ed.get("institution"), ed.get("end_date")] if x))
    if content.get("certifications"):
        d.add_heading("Certifications", level=1)
        for c in content["certifications"]:
            d.add_paragraph(c.get("name") if isinstance(c, dict) else str(c), style="List Bullet")
    if content.get("languages"):
        d.add_heading("Languages", level=1)
        d.add_paragraph(", ".join(
            (f"{lang.get('language')} ({lang.get('proficiency')})" if lang.get("proficiency") else lang.get("language"))
            if isinstance(lang, dict) else str(lang) for lang in content["languages"]))
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
