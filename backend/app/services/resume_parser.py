"""Resume text extraction (PDF / DOCX / TXT) and structured parsing (heuristic + optional AI)."""
from __future__ import annotations

import io
import re
from pathlib import Path

from app.ai.base import AIProvider, AIProviderError

SECTION_ALIASES = {
    "summary": ["summary", "professional summary", "profile", "about me", "objective", "career summary", "resumen", "perfil", "perfil profesional"],
    "experience": ["experience", "work experience", "professional experience", "employment history", "work history", "experiencia", "experiencia laboral", "experiencia profesional"],
    "education": ["education", "academic background", "educación", "educacion", "formación", "formacion"],
    "skills": ["skills", "core skills", "key skills", "technical skills", "competencies", "habilidades", "competencias", "core competencies"],
    "tools": ["tools", "software", "technologies", "tech stack", "herramientas"],
    "certifications": ["certifications", "certificates", "licenses", "certificaciones", "licenses & certifications"],
    "languages": ["languages", "idiomas"],
    "projects": ["projects", "proyectos", "selected projects"],
    "achievements": ["achievements", "accomplishments", "awards", "logros", "honors"],
}

MONTHS = "jan|feb|mar|apr|may|jun|jul|aug|sep|sept|oct|nov|dec|ene|abr|ago|dic"
DATE_RANGE_RE = re.compile(
    rf"(?P<start>(?:(?:{MONTHS})[a-z]*\.?\s+)?\d{{4}}|\d{{1,2}}/\d{{4}})\s*(?:-|–|—|to|a|hasta)\s*(?P<end>(?:(?:{MONTHS})[a-z]*\.?\s+)?\d{{4}}|\d{{1,2}}/\d{{4}}|present|current|now|actual|presente|today)",
    re.I,
)
MONTH_NUM = {m: i + 1 for i, m in enumerate(["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"])}
MONTH_NUM.update({"ene": 1, "abr": 4, "ago": 8, "dic": 12, "sept": 9})


class ResumeParseError(ValueError):
    pass


def extract_text(data: bytes, filename: str, mime: str | None = None) -> str:
    ext = Path(filename).suffix.lower()
    if ext == ".pdf" or mime == "application/pdf":
        from pypdf import PdfReader

        try:
            reader = PdfReader(io.BytesIO(data))
            text = "\n".join((page.extract_text() or "") for page in reader.pages)
        except Exception as e:  # pypdf raises many types
            raise ResumeParseError(f"Could not read PDF: {e}") from e
        if not text.strip():
            raise ResumeParseError("PDF has no extractable text (scanned image?). Export it as text-based PDF or DOCX.")
        return text
    if ext == ".docx":
        import docx

        try:
            d = docx.Document(io.BytesIO(data))
        except Exception as e:
            raise ResumeParseError(f"Could not read DOCX: {e}") from e
        lines = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                lines.append(" | ".join(c.text for c in row.cells))
        return "\n".join(lines)
    if ext in {".txt", ".md", ".text"} or (mime or "").startswith("text/"):
        for enc in ("utf-8", "latin-1"):
            try:
                return data.decode(enc)
            except UnicodeDecodeError:
                continue
    raise ResumeParseError(f"Unsupported file type '{ext}'. Use PDF, DOCX or TXT.")


def _norm_date(s: str | None) -> str | None:
    if not s:
        return None
    s = s.strip().lower()
    if s in {"present", "current", "now", "actual", "presente", "today"}:
        return "present"
    m = re.match(r"(\d{1,2})/(\d{4})", s)
    if m:
        return f"{m.group(2)}-{int(m.group(1)):02d}"
    m = re.match(rf"({MONTHS})[a-z]*\.?\s+(\d{{4}})", s)
    if m:
        return f"{m.group(2)}-{MONTH_NUM.get(m.group(1)[:4] if m.group(1).startswith('sept') else m.group(1)[:3], 1):02d}"
    m = re.match(r"(\d{4})", s)
    return m.group(1) if m else None


def _split_sections(lines: list[str]) -> dict[str, list[str]]:
    sections: dict[str, list[str]] = {"header": []}
    current = "header"
    for line in lines:
        key = line.strip().strip(":").strip().lower()
        found = None
        if 2 < len(key) < 40:
            for sec, aliases in SECTION_ALIASES.items():
                if key in aliases:
                    found = sec
                    break
        if found:
            current = found
            sections.setdefault(current, [])
        else:
            sections.setdefault(current, []).append(line)
    return sections


def _list_items(lines: list[str]) -> list[str]:
    items: list[str] = []
    for ln in lines:
        ln = ln.strip().lstrip("•-*·▪◦● ").strip()
        if not ln:
            continue
        if "," in ln or "|" in ln or ";" in ln:
            items.extend(x.strip() for x in re.split(r"[,|;•·]", ln) if x.strip())
        else:
            items.append(ln)
    clean = []
    for i in items:
        i = re.sub(r"^[A-Za-z /&]+:\s*", "", i) if ":" in i and len(i.split(":")[0]) < 25 else i
        if 1 < len(i) <= 60:
            clean.append(i)
    return list(dict.fromkeys(clean))


def _parse_experience(lines: list[str]) -> list[dict]:
    entries: list[dict] = []
    current: dict | None = None
    pending_header: list[str] = []
    for raw in lines:
        line = raw.strip()
        if not line:
            continue
        m = DATE_RANGE_RE.search(line)
        if m:
            if current:
                entries.append(current)
            head = (line[: m.start()] + line[m.end():]).strip(" |,-–—·")
            header_bits = [x.strip() for x in re.split(r"\s+(?:\||–|—|-|at|@|,)\s+", head) if x.strip()] if head else []
            if not header_bits and pending_header:
                header_bits = pending_header[-2:]
            title = header_bits[0] if header_bits else None
            company = header_bits[1] if len(header_bits) > 1 else None
            if not company and pending_header and title in pending_header:
                idx = pending_header.index(title)
                company = pending_header[idx + 1] if idx + 1 < len(pending_header) else None
            end = _norm_date(m.group("end"))
            current = {"title": title, "company": company, "location": None, "start_date": _norm_date(m.group("start")),
                       "end_date": None if end == "present" else end, "current": end == "present",
                       "description": "", "highlights": [], "skills": []}
            pending_header = []
        elif line[:1] in "•-*·▪◦●" or (current and len(line) > 60):
            if current:
                current["highlights"].append(line.lstrip("•-*·▪◦● ").strip())
        else:
            if current and current["highlights"]:
                entries.append(current)
                current = None
            if current and not current.get("company") and len(line) < 80:
                current["company"] = line
            elif current and len(line) >= 40:
                current["description"] = (current["description"] + " " + line).strip()
            else:
                pending_header.append(line)
    if current:
        entries.append(current)
    return [e for e in entries if e.get("title") or e.get("company")]


def _parse_education(lines: list[str]) -> list[dict]:
    out = []
    for ln in lines:
        ln = ln.strip().lstrip("•-* ")
        if not ln:
            continue
        m = DATE_RANGE_RE.search(ln)
        year = re.search(r"\b(19|20)\d{2}\b", ln)
        bits = [b.strip() for b in re.split(r"\s+[|–—-]\s+|,\s+", DATE_RANGE_RE.sub("", ln)) if b.strip()]
        if not bits:
            continue
        degree_first = re.search(r"bachelor|master|mba|phd|degree|diploma|licenciatura|ingenier|associate|b\.?s|b\.?a|certificate", bits[0], re.I)
        out.append({
            "institution": bits[1] if degree_first and len(bits) > 1 else bits[0],
            "degree": bits[0] if degree_first else (bits[1] if len(bits) > 1 else None),
            "field": None,
            "start_date": _norm_date(m.group("start")) if m else None,
            "end_date": _norm_date(m.group("end")) if m else (year.group(0) if year else None),
        })
    return out[:8]


def _parse_languages(lines: list[str]) -> list[dict]:
    out = []
    for item in _list_items(lines):
        m = re.match(r"([A-Za-zÀ-ÿ]+)\s*(?:[-–:(]\s*([A-Za-z0-9 ]+)\)?)?", item)
        if m:
            out.append({"language": m.group(1).strip(), "proficiency": (m.group(2) or "").strip() or None})
    return out


def heuristic_parse(text: str) -> dict:
    lines = [ln.rstrip() for ln in text.splitlines()]
    sections = _split_sections(lines)
    header = [ln.strip() for ln in sections.get("header", []) if ln.strip()]
    blob = text
    email = re.search(r"[\w.+-]+@[\w-]+\.[\w.-]+", blob)
    phone = re.search(r"(\+?\d[\d\s().-]{7,}\d)", blob)
    linkedin = re.search(r"(?:https?://)?(?:[a-z]{2,3}\.)?linkedin\.com/in/[\w\-%]+/?", blob, re.I)
    urls = re.findall(r"https?://[^\s,|)]+", blob)
    other_urls = [u for u in urls if "linkedin.com" not in u.lower()]
    name = None
    for ln in header[:4]:
        if not re.search(r"@|\d|http|linkedin", ln, re.I) and 1 < len(ln.split()) <= 5 and len(ln) < 50:
            name = ln.title() if ln.isupper() else ln
            break
    headline = None
    if name and header:
        idx = next((i for i, h in enumerate(header) if h.lower() == name.lower() or h == name), None)
        if idx is not None and idx + 1 < len(header) and not re.search(r"@|\d{3}|http", header[idx + 1]):
            headline = header[idx + 1]
    location = None
    for ln in header[:8]:
        m = re.search(r"([A-ZÀ-Ý][A-Za-zÀ-ÿ .]+,\s*[A-ZÀ-Ý][A-Za-zÀ-ÿ .]+)", ln)
        if m and "@" not in ln:
            location = m.group(1).strip()
            break
    summary = " ".join(x.strip() for x in sections.get("summary", []) if x.strip()) or None
    return {
        "full_name": name,
        "email": email.group(0) if email else None,
        "phone": phone.group(1).strip() if phone else None,
        "location": location,
        "linkedin_url": ("https://" + linkedin.group(0).split("://")[-1]) if linkedin else None,
        "portfolio_url": other_urls[0] if other_urls else None,
        "other_urls": other_urls[1:5],
        "headline": headline,
        "professional_summary": summary,
        "skills": _list_items(sections.get("skills", []))[:60],
        "tools": _list_items(sections.get("tools", []))[:40],
        "work_experience": _parse_experience(sections.get("experience", [])),
        "education": _parse_education(sections.get("education", [])),
        "certifications": [{"name": c, "issuer": None, "date": None} for c in _list_items(sections.get("certifications", []))][:20],
        "languages": _parse_languages(sections.get("languages", [])),
        "projects": [{"name": p[:80], "description": p, "url": None} for p in _list_items(sections.get("projects", []))][:10],
        "achievements": [a for a in (ln.strip().lstrip("•-* ") for ln in sections.get("achievements", [])) if a][:15],
        "_method": "heuristic",
    }


AI_SCHEMA_HINT = """{"full_name": str|null, "email": str|null, "phone": str|null, "location": str|null, "linkedin_url": str|null,
"portfolio_url": str|null, "other_urls": [str], "headline": str|null, "professional_summary": str|null, "skills": [str], "tools": [str],
"work_experience": [{"title": str, "company": str, "location": str|null, "start_date": "YYYY-MM"|null, "end_date": "YYYY-MM"|null,
"current": bool, "description": str, "highlights": [str], "skills": [str]}],
"education": [{"institution": str, "degree": str|null, "field": str|null, "start_date": str|null, "end_date": str|null}],
"certifications": [{"name": str, "issuer": str|null, "date": str|null}], "languages": [{"language": str, "proficiency": str|null}],
"projects": [{"name": str, "description": str, "url": str|null}], "achievements": [str], "industries": [str]}"""


def ai_parse(text: str, provider: AIProvider) -> dict | None:
    if not provider or not provider.is_available():
        return None
    try:
        data = provider.complete_json(
            "You extract structured data from resumes. Copy facts exactly as written. Never infer, guess or add "
            "information that is not explicitly in the resume text. Use null/empty lists for missing fields.",
            f"Return JSON with this shape:\n{AI_SCHEMA_HINT}\n\nRESUME TEXT:\n{text[:30000]}",
            max_tokens=8000,
        )
    except AIProviderError:
        return None
    if isinstance(data, dict):
        data["_method"] = f"ai:{provider.label}"
        return data
    return None


def parse_resume_text(text: str, provider: AIProvider | None = None) -> dict:
    return ai_parse(text, provider) or heuristic_parse(text)
